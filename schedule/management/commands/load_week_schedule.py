"""Разовая загрузка недельного расписания из файла вуза.

Команда запускается вручную из консоли, а не через сайт:

    python manage.py load_week_schedule "C:\\путь\\файл.xlsx" --replace

Веб-импорт удалён намеренно: он показывал «будет создано», а писал в
базу, и формат файла вуза ему не подходил. Эта команда делает то же
самое, но без двух опасных черт — её нельзя вызвать из браузера и она
сначала показывает, сколько записей собирается создать.

Файл вуза устроен сеткой: пять листов (1–4 курс и ОЗФ), в каждом строки
это дни недели, а группы стоят парами столбцов — «предмет» и «аудитория».
"""

import datetime
import re
from collections import Counter, defaultdict

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from schedule.exports import _pick_color
from schedule.models import (
    Group, Lesson, LessonTime, LessonType, Room, Subject, Substitution, Teacher,
)

# Тип занятия пишут в скобках. Ключи ищутся по началу, потому что в файле
# хватает вариантов «практ.зан», «практ.зан.» и «пакт.зан.».
TYPE_PATTERNS = [
    (re.compile(r'лабр|лаборатор', re.I), 'Лабораторная работа'),
    (re.compile(r'лекц\s*практ|лекцпракт', re.I), 'Лекция'),
    (re.compile(r'лекц|^лек\.', re.I), 'Лекция'),
    (re.compile(r'практ|пакт', re.I), 'Практика'),
    (re.compile(r'практикум', re.I), 'Практикум'),
    (re.compile(r'семинар', re.I), 'Семинар'),
    (re.compile(r'контрольн', re.I), 'Контрольная работа'),
]

# Должность отделяет преподавателя от предмета. Порядок важен:
# «ст. преп.» должен попасть в «Старший преподаватель», а не в
# «Преподаватель», а «доц» — раньше общего «преп».
POSITION_PATTERNS = [
    # В файле встречается и «ст. преп», и «ст перп.» — с пропущенной буквой.
    (re.compile(r'ст\s*\.?\s*(?:преп|перп)|старш\w*\s+преп', re.I), 'Старший преподаватель'),
    (re.compile(r'доц', re.I), 'Доцент'),
    (re.compile(r'проф', re.I), 'Профессор'),
    (re.compile(r'phd', re.I), 'PhD'),
    (re.compile(r'преп', re.I), 'Преподаватель'),
    (re.compile(r'ассистент', re.I), 'Ассистент'),
]

# Если должности в ячейке нет, преподавателя всё равно бывает видно по ФИО.
# Одно слово не подходит: им может быть и сам предмет.
NAME_ONLY = re.compile(
    r'^(?:'
    # «Саттарова Б.Ш.», «Архипова З.З.», «Крылова Г.Ш.»
    r'[А-ЯЁ][а-яё]+(?:\s+[А-ЯЁ][а-яё]+){0,2}(?:\s*[А-ЯЁ]\.){1,3}'
    # «Крылова Галина Сергеевна»
    r'|[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(?:\s+[А-ЯЁ][а-яё]+)?'
    r')$'
)

CURATOR = re.compile(r'кураторск\w*\s+час', re.I)
# Позиция, продублированная в файле: в «ст. ст. преп.» регулярное выражение
# находит «ст. преп.», и перед ФИО остаётся лишнее «ст. ». Точка
# обязательна, иначе от названия предмета откусило бы первые буквы.
LEFTOVER_POSITION = re.compile(r'\b(?:ст|доц|проф|асс|преп)\s*\.', re.I)
ROOM_NUM = re.compile(r'^(\d+)\s*/\s*(\d+)$')
SEMESTER = re.compile(r'на\s+([\d,\s]+?)\s*семестр', re.I)
WEEK_NO = re.compile(r'(\d{1,2})\s*недел', re.I)
CELL_DATE = re.compile(r'(\d{1,2})\.(\d{1,2})\.(\d{4})')


def _week_parity(ws):
    """Чётность недели по номеру недели из шапки или по дате из таблицы.

    Номер недели — прямое указание вуза, какой это нечётный или чётный счёт
    недель семестра. В шапке он есть не у всех файлов, поэтому вторым
    источником служит дата занятия: её считает та же функция, что и сайт.
    """
    from schedule.grid import _parity

    header = ' '.join(squash(ws.cell(r, c).value)
                      for r in range(1, 9) for c in range(1, 13))
    by_date = 0
    for r in range(1, 16):
        for c in range(1, 13):
            found = CELL_DATE.search(squash(ws.cell(r, c).value))
            if found:
                day = datetime.date(int(found.group(3)),
                                    int(found.group(2)),
                                    int(found.group(1)))
                by_date = _parity(day - datetime.timedelta(days=day.weekday()))
                break
        if by_date:
            break

    match = WEEK_NO.search(header)
    if not match:
        return by_date
    parity = 1 if int(match.group(1)) % 2 else 2
    if by_date and by_date != parity:
        # Молча выбрать одно из двух — значит показать занятия не в ту неделю.
        raise CommandError(
            'Не сходится чётность: в шапке %d неделя (это %s), '
            'а даты в таблице дают %s. Проверьте начало семестра в настройках.'
            % (int(match.group(1)),
               'нечётная' if parity == 1 else 'чётная',
               'нечётную' if by_date == 1 else 'чётную'))
    return parity


def parity_of(ws):
    """Чётность недель по шапке листа.

    Сначала номер недели из шапки («8 неделя с 19 октября по 23 октября»):
    именно он отличает чётную неделю от нечётной, потому что строка
    «на 1 семестр» одинакова в файлах обеих недель и по ней не различить.
    Потом дата из таблицы. И только если нет ни того, ни другого — семестр:
    вуз пишет «на 1 семестр», «на 3 семестр», «на 5 семестр»,
    «на 7 семестр» и «на 1,3,5,7,9 семестр» для ОЗФ. Если в шапке написаны
    и чётные, и нечётные семестры сразу, недель не различить — тогда 0.
    """
    week_parity = _week_parity(ws)
    if week_parity:
        return week_parity

    header = ' '.join(squash(ws.cell(r, c).value)
                      for r in range(1, 4) for c in range(1, 9))
    match = SEMESTER.search(header)
    if not match:
        return 0
    numbers = [int(n) for n in re.findall(r'\d+', match.group(1))]
    if not numbers:
        return 0
    if all(n % 2 for n in numbers):
        return 1
    if all(n % 2 == 0 for n in numbers):
        return 2
    return 0


def squash(text):
    """Одна строка без лишних пробелов: в файле текст разбит переносами."""
    return re.sub(r'\s+', ' ', str(text or '').replace('\u00a0', ' ')).strip()


def tidy_name(text):
    """«Соколова С» -> «Соколова С.», «Машарипов С.И.» не трогаем.

    Заодно чинит то, что осталось после отделения должности: там почти
    всегда остаётся точка («ст. преп.»), а в файле есть и «АрхиповаЗ.З.»,
    где между фамилией и инициалами не хватает пробела.
    """
    value = squash(text).lstrip(' .,;:')
    value = re.sub(r'\.{2,}', '.', value)
    value = re.sub(r'(?<=[а-яё])(?=[А-ЯЁ]\.)', ' ', value)
    parts = []
    for part in value.split(' '):
        if re.fullmatch(r'[А-ЯЁ]|[А-ЯЁ](?:\.[А-ЯЁ])+', part):
            part += '.'
        parts.append(part)
    return ' '.join(parts)


def split_cell(text):
    """Ячейка занятия -> (предмет, тип, преподаватель, должность)."""
    rest = squash(text)
    # В файле есть склейка: «бизнеселекц.)» вместо «бизнесе (лекц.)».
    # Без пробела скобка не открывается и тип не виден.
    rest = re.sub(r'(?<=[а-яё])(?=(?:лекц|практ|пакт|лабр)[^)]*\))', ' (', rest)

    # Тип ищем во всех скобках подряд, а не только в первой: встречается
    # «История экономики (на английском языке) (лекц.)», где первая скобка
    # к типу отношения не имеет.
    lesson_type = None
    for match in re.finditer(r'\(([^)]{2,40})\)', rest):
        inner = match.group(1)
        for pattern, name in TYPE_PATTERNS:
            if pattern.search(inner):
                lesson_type = name
                rest = rest[:match.start()] + ' ' + rest[match.end():]
                break
        if lesson_type:
            break

    teacher = position = None
    for pattern, name in POSITION_PATTERNS:
        matches = list(pattern.finditer(rest))
        if not matches:
            continue
        position = name
        # ФИО — это всё, что после должности. Раньше бралась одна позиция,
        # теперь их может быть несколько («ст. ст. преп.»), поэтому ФИО
        # берётся после последней вхождения, а предмет заканчивается перед
        # первым. Так «доц.» уходит из названия вместе со своей точкой.
        teacher = tidy_name(rest[matches[-1].end():])
        rest = rest[:matches[0].start()]
        break
    else:
        # Должности нет — ФИО ищем в конце ячейки. ФИО занимает
        # от одного до трёх слов, поэтому перебираем хвосты подлиннее.
        words = rest.split(' ')
        for take in (3, 2, 1):
            if len(words) <= take:
                continue
            tail = ' '.join(words[-take:])
            if NAME_ONLY.match(tail):
                teacher = tidy_name(tail)
                rest = ' '.join(words[:-take])
                break

    # Остаток продублированной позиции: «ст. » перед ФИО, когда полное
    # сочетание «ст. преп.» уже вырезано. Чистим только если должность
    # действительно нашлась: без этого сработала бы лишняя точка в
    # предмете вида «Основы предпринимательства ст.».
    if position:
        rest = LEFTOVER_POSITION.sub(' ', rest)

    subject = re.sub(r'\s+', ' ', rest).strip(' .,;:-–—')
    return subject or None, lesson_type, (teacher or None), position


def parse_room(text):
    """«3/308» -> корпус 3, аудитория 308. «спорт зал» -> как есть."""
    value = squash(text)
    match = ROOM_NUM.match(value)
    if match:
        return match.group(2), match.group(1)
    return value or None, None


def course_of(code):
    """«6301» -> 1 курс, «3202» -> 4 курс: первая цифра кода — это 7 минус курс."""
    return 7 - int(code[0])


def parse_time(value):
    """«9.00-10.20» и «15:30-16:50» -> (time, time)."""
    match = re.match(r'(\d{1,2})[.:](\d{2})\s*[-–—]\s*(\d{1,2})[.:](\d{2})', squash(value))
    if not match:
        return None
    start = datetime.time(int(match[1]), int(match[2]))
    end = datetime.time(int(match[3]), int(match[4]))
    return start, end


def find_columns(ws):
    """Колонки «Дата / День / Пара / Время» по заголовкам."""
    found = {}
    for r in range(1, 12):
        for c in range(1, 16):
            head = squash(ws.cell(r, c).value).lower()
            if head.startswith('дата'):
                found['date'] = c
            elif head.startswith(('дни', 'день')):
                found['day'] = c
            elif head.startswith('пара'):
                found['pair'] = c
            elif head.startswith('время'):
                found['time'] = c
    return found


def find_groups(ws):
    """Столбцы групп: код вида 6301 и «ауд» рядом + профиль.

    Проверяется и код, и соседняя колонка с «ауд»: иначе за группу можно
    принять «2026-2027 уч.» из заголовка, а такая группа одна — значит,
    лист разобрался бы неправильно и тихо.
    """
    for r in range(1, 14):
        codes = [c for c in range(1, ws.max_column + 1)
                 if re.match(r'^\d{4}\b', squash(ws.cell(r, c).value))
                 and squash(ws.cell(r, c + 1).value).lower().startswith('ауд')]
        if codes:
            profiles = {}
            head_row = r - 1
            for c in codes:
                # Объединение живёт на строке выше кода группы, а не на самой
                # строке кодов: иначе профиль остаётся пустым.
                value = ''
                for rng in ws.merged_cells.ranges:
                    if (rng.min_row <= head_row <= rng.max_row
                            and rng.min_col <= c <= rng.max_col):
                        value = squash(ws.cell(rng.min_row, rng.min_col).value)
                        break
                if not value:
                    value = squash(ws.cell(head_row, c).value)
                value = re.sub(r'^профиль\w*\s*:?\s*', '', value, flags=re.I)
                profiles[c] = value
            return r, codes, profiles
    return None, [], {}


def read_sheet(ws, sheet_course=None):
    """Лист -> список записей и статистика."""
    cols = find_columns(ws)
    group_row, group_cols, profiles = find_groups(ws)
    if not group_cols:
        raise CommandError(f'На листе «{ws.title.strip()}» не нашлись группы.')

    last = max(r for r in range(1, ws.max_row + 1)
               if any(squash(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)))

    day = None
    records = []
    skipped = Counter()
    parity = parity_of(ws)
    for r in range(group_row + 1, last + 1):
        raw_date = squash(ws.cell(group_row and r, cols['date']).value) if 'date' in cols else ''
        if 'date' in cols and raw_date:
            parsed = re.match(r'(\d{1,2})\.(\d{2})\.(\d{4})', raw_date)
            if parsed:
                date = datetime.date(int(parsed[3]), int(parsed[2]), int(parsed[1]))
                day = date.isoweekday()
        if not day:
            continue
        pair = ws.cell(r, cols['pair']).value if 'pair' in cols else None
        pair = int(pair) if isinstance(pair, (int, float)) or squash(pair).isdigit() else None
        time = parse_time(ws.cell(r, cols['time']).value) if 'time' in cols else None

        for col in group_cols:
            raw = ws.cell(r, col).value
            text = squash(raw)
            if not text or text in ('.', '-', '—'):
                continue
            if CURATOR.search(text):
                skipped['кураторский час'] += 1
                continue
            if not pair or not time:
                skipped['без номера пары или времени'] += 1
                continue

            code = squash(ws.cell(group_row, col).value)
            match = re.match(r'^(\d{4})(?:-(\d{4}))?', code)
            name = match[1] if match else code
            subject, lesson_type, teacher, position = split_cell(text)
            if not subject:
                skipped['пустой предмет'] += 1
                continue
            room_name, building = parse_room(ws.cell(r, col + 1).value)
            records.append({
                'group': name,
                'course': course_of(name),
                'specialty': profiles.get(col, ''),
                'day': day,
                'pair': pair,
                'parity': parity,
                'time': time,
                'subject': subject,
                'lesson_type': lesson_type,
                'teacher': teacher,
                'position': position,
                'room': room_name,
                'building': building,
                'raw': text,
                'sheet': ws.title.strip(),
            })
    return records, skipped


def read_file(path):
    """Книга -> (записи, пропущенное, время пар)."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except FileNotFoundError:
        raise CommandError(f'Файл не найден: {path}')
    except Exception as exc:
        raise CommandError(f'Не удалось прочитать файл: {exc}')

    records, skipped = [], Counter()
    for ws in wb.worksheets:
        part, part_skipped = read_sheet(ws)
        records.extend(part)
        skipped.update(part_skipped)

    # Время пары берём тем, что встречается чаще: в среду и четверг в файле
    # сдвиг на 5 минут, а «пара» в системе одна на неделю.
    votes = defaultdict(Counter)
    for rec in records:
        votes[rec['pair']][rec['time']] += 1
    pair_times = {number: counter.most_common(1)[0][0]
                  for number, counter in votes.items()}
    return records, skipped, pair_times


class Command(BaseCommand):
    help = 'Загрузить недельное расписание из файла вуза (запуск из консоли).'

    def add_arguments(self, parser):
        parser.add_argument('path', help='путь к .xlsx с расписанием')
        parser.add_argument(
            '--replace', action='store_true',
            help='очистить прежнее расписание и справочники перед загрузкой',
        )
        parser.add_argument(
            '--noinput', action='store_true',
            help='не спрашивать подтверждение перед очисткой (--replace)',
        )
        parser.add_argument(
            '--dry-run', action='store_true',
            help='только показать, что будет загружено, ничего не записывая',
        )

    def handle(self, *args, **options):
        records, skipped, pair_times = read_file(options['path'])
        if not records:
            raise CommandError('Не нашлось ни одного занятия — проверьте файл.')

        self.stdout.write(f'  занятий найдено: {len(records)}')
        self.stdout.write(f'  групп: {len({r["group"] for r in records})}')
        self.stdout.write(f'  предметов: {len({r["subject"] for r in records})}')
        self.stdout.write(f'  преподавателей: {len({r["teacher"] for r in records if r["teacher"]})}')
        self.stdout.write(f'  аудиторий: {len({r["room"] for r in records if r["room"]})}')

        parity_names = {0: 'каждую неделю', 1: 'нечётные недели', 2: 'чётные недели'}
        for parity, count in sorted(Counter(r['parity'] for r in records).items()):
            self.stdout.write(f'  недель: {parity_names[parity]} — {count}')
        for reason, count in skipped.items():
            self.stdout.write(f'  пропущено ({reason}): {count}')

        odd = Counter(r['time'] for r in records if r['time'] != pair_times[r['pair']])
        if odd:
            self.stdout.write(self.style.WARNING(
                f'  внимание: время пары в файле местами отличается, взято самое частое. '
                f'Отличия: {len(odd)}'
            ))

        if options['dry_run']:
            self.stdout.write('  пробный запуск: база не менялась.')
            return

        if options['replace'] and not options['noinput'] and not options['dry_run']:
            # Здесь удаляется всё прежнее расписание и справочники, поэтому
            # вопрос обязателен: «--replace» в скрипте может стоять рядом с
            # чужим кодом, и молчаливая очистка базы — плохая идея.
            self.stdout.write(self.style.WARNING(
                '  ВНИМАНИЕ: прежнее расписание и справочники будут удалены.'
                ' Сделайте копию db.sqlite3.'
            ))
            answer = input('  Точно очистить и загрузить заново? (да/нет) ')
            if answer.strip().lower() not in ('да', 'y', 'yes'):
                self.stdout.write('  Отменено.')
                return

        with transaction.atomic():
            if options['replace']:
                self._clear()
            count = self._load(records, pair_times)

        self.stdout.write(self.style.SUCCESS(f'  записано занятий: {count}'))

    def _clear(self):
        Substitution.objects.all().delete()
        Lesson.objects.all().delete()
        # Предмет и пара защищены от удаления, пока на них ссылаются
        # занятия, поэтому справочники чистим после них.
        Subject.objects.all().delete()
        LessonType.objects.all().delete()
        Room.objects.all().delete()
        Teacher.objects.all().delete()
        Group.objects.all().delete()
        self.stdout.write('  прежнее расписание и демо-справочники удалены.')

    def _load(self, records, pair_times):
        groups, subjects, types, teachers, rooms, times = {}, {}, {}, {}, {}, {}

        def one(cache, model, key, lookup, **fields):
            """Найти или создать запись справочника по нормализованному ключу."""
            obj = cache.get(key)
            if obj is None:
                obj = model.objects.filter(**lookup).first()
                if obj is None:
                    obj = model.objects.create(**fields)
                    obj.fill_norm()
                    obj.save()
                cache[key] = obj
            return obj

        for number, (start, end) in sorted(pair_times.items()):
            obj, _ = LessonTime.objects.update_or_create(
                number=number, defaults={'start': start, 'end': end})
            times[number] = obj

        lessons = []
        for rec in records:
            group = one(
                groups, Group, (rec['group'], rec['course']),
                Group.norm_lookup(rec['group'], rec['course']),
                name=rec['group'], course=rec['course'],
                specialty=rec['specialty'],
            )
            subject = one(
                subjects, Subject, rec['subject'],
                Subject.norm_lookup(rec['subject']),
                name=rec['subject'], color=_pick_color(rec['subject']),
            )
            lesson_type = None
            if rec['lesson_type']:
                lesson_type = one(
                    types, LessonType, rec['lesson_type'],
                    LessonType.norm_lookup(rec['lesson_type']),
                    name=rec['lesson_type'],
                )
            teacher = None
            if rec['teacher']:
                teacher = one(
                    teachers, Teacher, rec['teacher'],
                    Teacher.norm_lookup(rec['teacher']),
                    full_name=rec['teacher'], position=rec['position'] or '',
                )
            room = None
            if rec['room']:
                room = one(
                    rooms, Room, (rec['room'], rec['building']),
                    Room.norm_lookup(rec['room'], rec['building'] or ''),
                    name=rec['room'], building=rec['building'] or '',
                )
            lessons.append(Lesson(
                group=group, subject=subject, lesson_type=lesson_type,
                teacher=teacher, room=room, day=rec['day'],
                lesson_time=times[rec['pair']], week_parity=rec['parity'],
            ))

        # Повторный запуск без --replace не должен падать: в базе уже есть
        # ровно такое занятие, и ограничение уникальности его отвергнет.
        # Такие строки пропускаем и считаем отдельно.
        existing = set(Lesson.objects.filter(
            group_id__in={group.pk for group in groups.values()}
        ).values_list('group_id', 'day', 'lesson_time_id', 'week_parity', 'subject_id'))

        fresh = []
        skipped = 0
        for lesson in lessons:
            key = (lesson.group_id, lesson.day, lesson.lesson_time_id,
                   lesson.week_parity, lesson.subject_id)
            if key in existing:
                skipped += 1
                continue
            existing.add(key)
            fresh.append(lesson)

        Lesson.objects.bulk_create(fresh)
        if skipped:
            self.stdout.write(f'  уже было в базе, пропущено: {skipped}')
        return len(fresh)